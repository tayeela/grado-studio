"""PPTX-бэкенд: тот же список операций, что у PDF и SVG, — в редактируемую презентацию PowerPoint.

Операция становится фигурой: контуры и линии — полилиниями (custGeom), подписи — надписями с настоящим текстом,
снимки и тайлы подложки — рисунками. Фигуры одной рамки листа собраны в группу (карта, легенда, оформление),
внутри карты — по слоям (операция "layer"). Обрезки в PowerPoint нет, поэтому она делается геометрией: контуры
режутся по прямоугольнику рамки, рисунки кадрируются, подписи и значки вне рамки пропускаются, а обрезка по
контуру (округ на схеме шаблона ППТ КРТ) — белой маской поверх: прямоугольник с вырезом по контуру.

Пакет PPTX (шаблон, части, связи, заметки) собирает python-pptx (MIT); фигуры пишутся разметкой DrawingML и
вставляются в дерево слайда, шрифты встраиваются частями его же пакета. Зовёт apps/grado2/print_io.py (render_pptx);
тесты — core/render/tests/test_pptx.py, apps/grado2/tests/test_t86_pptx_speed.py.
"""

import hashlib
import struct
import io
import math
import re
from pathlib import Path
from xml.sax.saxutils import escape

from .ttf import Font, find_font

EMU = 36000                                          # EMU в мм
# семейство операции → гарнитура PowerPoint и жирность (те же шрифты, что встраивает PDF)
FONTS = {"ui": ("PT Sans", False), "ui-bold": ("PT Sans", True), "serif": ("PT Serif", False),
         "serif-bold": ("PT Serif", True), "sans": ("Montserrat", False), "sans-bold": ("Montserrat", True),
         "mono": ("Courier New", False)}
NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')
XML = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _hex(c):
    return "%02X%02X%02X" % tuple(max(0, min(255, int(round(float(v) * 255)))) for v in c)


def _t(s):
    return escape(str(s), {'"': "&quot;"})


# ---------- геометрия обрезки ----------

def clip_polygon(ring, box):
    """Кольцо по прямоугольнику box = (x0, y0, x1, y1), Сазерленд — Ходжмен."""
    x0, y0, x1, y1 = box
    pts = list(ring)
    if len(pts) > 1 and pts[0] == pts[-1]:
        pts = pts[:-1]
    ix = lambda x: lambda a, b: (x, a[1] + (b[1] - a[1]) * (x - a[0]) / (b[0] - a[0]))
    iy = lambda y: lambda a, b: (a[0] + (b[0] - a[0]) * (y - a[1]) / (b[1] - a[1]), y)
    for inside, cut in ((lambda p: p[0] >= x0, ix(x0)), (lambda p: p[0] <= x1, ix(x1)),
                        (lambda p: p[1] >= y0, iy(y0)), (lambda p: p[1] <= y1, iy(y1))):
        out = []
        for i, cur in enumerate(pts):
            prev = pts[i - 1]
            if inside(cur):
                if not inside(prev):
                    out.append(cut(prev, cur))
                out.append(cur)
            elif inside(prev):
                out.append(cut(prev, cur))
        pts = out
        if not pts:
            break
    return pts


def _segment(a, b, box):
    """Отрезок по прямоугольнику (Лян — Барски): концы внутри или None."""
    x0, y0, x1, y1 = box
    dx, dy = b[0] - a[0], b[1] - a[1]
    t0, t1 = 0.0, 1.0
    for p, q in ((-dx, a[0] - x0), (dx, x1 - a[0]), (-dy, a[1] - y0), (dy, y1 - a[1])):
        if p == 0:
            if q < 0:
                return None
        elif p < 0:
            t0 = max(t0, q / p)
        else:
            t1 = min(t1, q / p)
        if t0 > t1:
            return None
    return (a if t0 == 0 else (a[0] + t0 * dx, a[1] + t0 * dy)), (b if t1 == 1 else (a[0] + t1 * dx, a[1] + t1 * dy))


def clip_polyline(pts, box):
    """Ломаная по прямоугольнику: куски внутри."""
    runs, cur = [], []
    for a, b in zip(pts, pts[1:]):
        seg = _segment(a, b, box)
        if seg is None:
            if len(cur) > 1:
                runs.append(cur)
            cur = []
            continue
        p, q = seg
        if cur and cur[-1] == p:
            cur.append(q)
        else:
            if len(cur) > 1:
                runs.append(cur)
            cur = [p, q]
    if len(cur) > 1:
        runs.append(cur)
    return runs


def _bbox(rings):
    xs = [x for r in rings for x, _ in r]
    ys = [y for r in rings for _, y in r]
    return min(xs), min(ys), max(xs), max(ys)


def _inside_box(bb, box):
    return bb[0] >= box[0] and bb[1] >= box[1] and bb[2] <= box[2] and bb[3] <= box[3]


def _meet(a, b):
    if a is None:
        return b
    return max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3])


def _area2(r):
    return sum(x1 * y2 - x2 * y1 for (x1, y1), (x2, y2) in zip(r, r[1:] + r[:1]))


def _in_ring(pt, r):
    x, y, inside = pt[0], pt[1], False
    for (x1, y1), (x2, y2) in zip(r, r[1:] + r[:1]):
        if (y1 > y) != (y2 > y) and x < x1 + (y - y1) * (x2 - x1) / (y2 - y1):
            inside = not inside
    return inside


def orient(rings):
    """Кольца с дырами: наружные — в одну сторону обхода, дыры — в другую (по глубине вложенности). Тогда дыра
    остаётся пустой при любом правиле заливки — и по чётности (PDF), и по ненулевому числу оборотов."""
    if len(rings) < 2 or len(rings) > 300:
        return rings
    out = []
    for i, r in enumerate(rings):
        probe = ((r[0][0] + r[1][0]) / 2, (r[0][1] + r[1][1]) / 2) if len(r) > 1 else r[0]
        depth = sum(1 for j, o in enumerate(rings) if j != i and len(o) > 2 and _in_ring(probe, o))
        out.append(r if (_area2(r) > 0) == (depth % 2 == 0) else r[::-1])
    return out


# ---------- фигуры ----------

def _eot(ttf):
    """TrueType → EOT (Embedded OpenType, версия 2.1, без сжатия и шифрования): так PowerPoint хранит ppt/fonts/*.fntdata."""
    t = {ttf[o:o + 4].decode("latin1"): struct.unpack(">II", ttf[o + 8:o + 16]) for o in
         (12 + 16 * i for i in range(struct.unpack(">H", ttf[4:6])[0]))}
    sub = lambda tag: ttf[t[tag][0]:t[tag][0] + t[tag][1]]
    os2, head, name = sub("OS/2"), sub("head"), sub("name")
    ver, = struct.unpack(">H", os2[:2])
    weight, fstype = struct.unpack(">HH", os2[4:6] + os2[8:10])
    panose = os2[32:42]
    ur = struct.unpack(">IIII", os2[42:58])
    cp = struct.unpack(">II", os2[78:86]) if ver >= 1 else (1, 0)
    italic = 1 if struct.unpack(">H", os2[62:64])[0] & 1 else 0
    cnt, off = struct.unpack(">HH", name[2:6])
    names = {}
    for i in range(cnt):
        pid, eid, lid, nid, ln, so = struct.unpack(">6H", name[6 + 12 * i:18 + 12 * i])
        raw = name[off + so:off + so + ln]
        if pid == 3 and eid in (1, 10) and nid in (1, 2, 5, 4):
            names.setdefault(nid, raw.decode("utf-16-be").encode("utf-16-le"))
        elif pid == 1 and nid in (1, 2, 5, 4):
            names.setdefault(nid, raw.decode("mac_roman", "replace").encode("utf-16-le"))
    strs = b"".join(struct.pack("<HH", 0, len(names.get(n, b""))) + names.get(n, b"") for n in (1, 2, 5, 4)) + struct.pack("<HH", 0, 0)
    body = (struct.pack("<II", 0x00020001, 0) + panose + struct.pack("<BBIHH", 1, italic, weight, fstype, 0x504C)
            + struct.pack("<IIII", *ur) + struct.pack("<II", *cp) + struct.pack("<I", struct.unpack(">I", head[8:12])[0])
            + b"\0" * 16 + strs)
    return struct.pack("<II", 8 + len(body) + len(ttf), len(ttf)) + body + ttf


class _Slide:
    def __init__(self, pres, k):
        self.pres, self.k = pres, k
        self.ids, self.rels = 1, {}                       # rels: имя картинки → rId
        self.stack = [[]]                                  # [[(xml, bbox)]] — открытые группы
        self.names = [None]
        self.clips = []
        self.pending = None                                # подряд идущие линии одного вида — одной фигурой
        self.alpha = None                                  # прозрачность заливки текущей фигуры

    def e(self, v):
        # float(): числовой атрибут разметки не должен пропустить строку (v * int повторяет её)
        return int(round(float(v) * EMU * self.k))

    def nid(self):
        self.ids += 1
        return self.ids

    def add(self, xml, bb, direct=False):
        if not direct:
            self.flush()
        self.stack[-1].append((xml, bb))

    def flush(self):
        """Линии без заливки одного цвета, толщины и штриха, идущие подряд (штриховка, сетка, контуры одного слоя),
        сливаются в одну фигуру: на листе с зонами иначе выходят десятки тысяч фигур."""
        if self.pending:
            (closed, stroke, width, dash, join, name), rings = self.pending
            self.pending = None
            self._emit(rings, closed, None, stroke, width, dash, join, name)

    # группы: рамка листа → слои карты
    def open(self, name):
        self.flush()
        self.stack.append([])
        self.names.append(name)

    def close(self):
        self.flush()
        items, name = self.stack.pop(), self.names.pop()
        if not items:
            return
        if len(items) == 1:                                       # одна фигура — без группы, под именем слоя
            xml, bb = items[0]
            if name:
                xml = re.sub(r'(<p:cNvPr id="\d+" name=")[^"]*"', lambda m: m.group(1) + _t(name) + '"', xml, count=1)
            self.add(xml, bb)
            return
        bb = (min(b[0] for _, b in items), min(b[1] for _, b in items), max(b[2] for _, b in items), max(b[3] for _, b in items))
        x, y, cx, cy = bb[0], bb[1], max(1, bb[2] - bb[0]), max(1, bb[3] - bb[1])
        xf = f'<a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/><a:chOff x="{x}" y="{y}"/><a:chExt cx="{cx}" cy="{cy}"/>'
        self.add(f'<p:grpSp><p:nvGrpSpPr><p:cNvPr id="{self.nid()}" name="{_t(name or "Группа")}"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
                 f'<p:grpSpPr><a:xfrm>{xf}</a:xfrm></p:grpSpPr>{"".join(x for x, _ in items)}</p:grpSp>', bb)

    def box(self):
        box = None
        for c in self.clips:
            box = _meet(box, c["box"])
        return box

    # --- контур ---
    def path(self, rings, closed, fill, stroke, width, dash, join, name="Контур"):
        rings = [r for r in rings if len(r) >= 2]
        if not rings or (fill is None and stroke is None):
            return
        box = self.box()
        if box is not None:
            bb = _bbox(rings)
            if bb[2] < box[0] or bb[0] > box[2] or bb[3] < box[1] or bb[1] > box[3]:
                return
            if not _inside_box(bb, box):
                if closed and fill is not None:
                    cut = [c for c in (clip_polygon(r, box) for r in rings) if len(c) >= 3]
                    if cut:
                        self._shape(orient(cut), True, fill, None, width, dash, join, name)
                if stroke is not None:
                    runs = [run for r in rings for run in clip_polyline(list(r) + ([r[0]] if closed else []), box)]
                    if runs:
                        self._shape(runs, False, None, stroke, width, dash, join, name)
                return
        self._shape(orient(rings) if closed and fill is not None else rings, closed, fill, stroke, width, dash, join, name)

    def _shape(self, rings, closed, fill, stroke, width, dash, join, name):
        if fill is None:
            key = (bool(closed), tuple(stroke), width, tuple(dash) if dash else None, join, name)
            if self.pending and self.pending[0] == key and sum(len(r) for r in self.pending[1]) < 20000:
                self.pending[1].extend(rings)
                return
            self.flush()
            self.pending = (key, list(rings))
            return
        self.flush()
        self._emit(rings, closed, fill, stroke, width, dash, join, name)

    def _emit(self, rings, closed, fill, stroke, width, dash, join, name):
        pts = [[(self.e(x), self.e(y)) for x, y in r] for r in rings]
        x0 = min(x for r in pts for x, _ in r)
        y0 = min(y for r in pts for _, y in r)
        cx = max(1, max(x for r in pts for x, _ in r) - x0)
        cy = max(1, max(y for r in pts for _, y in r) - y0)
        geo = "".join(f'<a:moveTo><a:pt x="{r[0][0] - x0}" y="{r[0][1] - y0}"/></a:moveTo>'
                      + "".join(f'<a:lnTo><a:pt x="{x - x0}" y="{y - y0}"/></a:lnTo>' for x, y in r[1:])
                      + ("<a:close/>" if closed else "") for r in pts)
        nofill = "" if fill is not None else ' fill="none"'
        alpha = f'<a:alpha val="{int(float(self.alpha) * 100000)}"/>' if fill is not None and self.alpha is not None and float(self.alpha) < 1 else ""
        clr = f'<a:srgbClr val="{_hex(fill)}">{alpha}</a:srgbClr>' if alpha else f'<a:srgbClr val="{_hex(fill)}"/>' if fill is not None else ""
        fill_xml = f'<a:solidFill>{clr}</a:solidFill>' if fill is not None else "<a:noFill/>"
        self.add(f'<p:sp><p:nvSpPr><p:cNvPr id="{self.nid()}" name="{_t(name)}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
                 f'<a:xfrm><a:off x="{x0}" y="{y0}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
                 f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="l" t="t" r="r" b="b"/>'
                 f'<a:pathLst><a:path w="{cx}" h="{cy}"{nofill}>{geo}</a:path></a:pathLst></a:custGeom>'
                 f'{fill_xml}{self._line(stroke, width, dash, join)}</p:spPr></p:sp>',
                 (x0, y0, x0 + cx, y0 + cy), direct=True)

    def _line(self, stroke, width, dash, join):
        if stroke is None:
            return "<a:ln><a:noFill/></a:ln>"
        w = max(0.05, float(width or 0))
        d = ""
        if dash:
            ds = [float(v) for v in dash] * (2 if len(dash) % 2 else 1)
            d = "<a:custDash>" + "".join(f'<a:ds d="{max(1, int(ds[i] / w * 100000))}" sp="{max(1, int(ds[i + 1] / w * 100000))}"/>'
                                         for i in range(0, len(ds), 2)) + "</a:custDash>"
        j = "<a:round/>" if join == "round" else '<a:miter lim="1000000"/>'
        return f'<a:ln w="{self.e(w)}" cap="flat"><a:solidFill><a:srgbClr val="{_hex(stroke)}"/></a:solidFill>{d}{j}</a:ln>'

    # --- значок ---
    def circle(self, o):
        box = self.box()
        if box is not None and not (box[0] <= o["x"] <= box[2] and box[1] <= o["y"] <= box[3]):
            return
        r = o["r"]
        x, y, d = self.e(o["x"] - r), self.e(o["y"] - r), max(1, self.e(2 * r))
        fill = f'<a:solidFill><a:srgbClr val="{_hex(o["fill"])}"/></a:solidFill>' if o.get("fill") is not None else "<a:noFill/>"
        self.add(f'<p:sp><p:nvSpPr><p:cNvPr id="{self.nid()}" name="Значок"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
                 f'<a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{d}" cy="{d}"/></a:xfrm><a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom>'
                 f'{fill}{self._line(o.get("stroke"), o.get("width", 0.3), None, "round")}</p:spPr></p:sp>', (x, y, x + d, y + d))

    # --- надпись ---
    def text(self, o):
        box = self.box()
        if box is not None and not (box[0] <= o["x"] <= box[2] and box[1] <= o["y"] <= box[3]):
            return
        fam = o.get("font_family") or "ui"
        fam = fam if fam in FONTS else "ui"
        font = self.pres.font(fam)
        size = float(o["size"])
        asc, dsc = font.ascent / font.units_per_em * size, -font.descent / font.units_per_em * size
        w, h = font.text_width(o["text"], size) * 1.04 + size * .2, asc + dsc
        anchor = o.get("anchor")
        lx = -w / 2 if anchor == "middle" else -w if anchor == "end" else -size * .1
        ang = o.get("angle", 0.0) or 0.0
        cx, cy = lx + w / 2, -asc + h / 2                          # середина надписи от точки привязки (на базовой линии)
        ca, sa = math.cos(ang), math.sin(ang)
        mx, my = o["x"] + cx * ca - cy * sa, o["y"] + cx * sa + cy * ca
        x, y, ew, eh = self.e(mx - w / 2), self.e(my - h / 2), max(1, self.e(w)), max(1, self.e(h))
        rot = f' rot="{int(round(math.degrees(ang) * 60000)) % 21600000}"' if ang else ""
        face, bold = FONTS[fam]
        algn = {"middle": "ctr", "end": "r"}.get(anchor, "l")
        glow = ""
        if o.get("halo"):                                          # ореол подписи — «свечение», как в шаблоне владельца
            glow = (f'<a:effectLst><a:glow rad="{self.e(float(o["halo"]) * 1.6)}">'
                    f'<a:srgbClr val="{_hex(o.get("halo_color", (1, 1, 1)))}"/></a:glow></a:effectLst>')
        sz = max(100, int(round(size / 0.3528 * self.k * 100)))
        b = (' b="1"' if bold else "") + (' u="sng"' if o.get("underline") else "")
        self.add(f'<p:sp><p:nvSpPr><p:cNvPr id="{self.nid()}" name="Надпись"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr>'
                 f'<a:xfrm{rot}><a:off x="{x}" y="{y}"/><a:ext cx="{ew}" cy="{eh}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
                 f'<a:noFill/></p:spPr><p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="t">'
                 f'<a:noAutofit/></a:bodyPr><a:lstStyle/><a:p><a:pPr algn="{algn}"/><a:r><a:rPr lang="ru-RU" sz="{sz}"'
                 f'{b} dirty="0"><a:solidFill><a:srgbClr val="{_hex(o.get("color", (0, 0, 0)))}"/></a:solidFill>'
                 f'{glow}<a:latin typeface="{face}"/><a:cs typeface="{face}"/></a:rPr><a:t>{_t(o["text"])}</a:t></a:r></a:p></p:txBody></p:sp>',
                 (x, y, x + ew, y + eh))

    # --- рисунок ---
    def image(self, o):
        key = self.pres.media(o["data"])
        rid = self.rels.setdefault(key, f"@img{len(self.rels) + 1}@")    # настоящий rId даст python-pptx при сборке
        box, crop, rot = self.box(), "", ""
        if "quad" in o:                                            # тайл параллелограммом: поворот и растяжение
            (tlx, tly), (trx, try_), (blx, bly) = o["quad"]
            w, h = math.hypot(trx - tlx, try_ - tly), math.hypot(blx - tlx, bly - tly)
            cx, cy = (trx + blx) / 2, (try_ + bly) / 2
            ang = math.atan2(try_ - tly, trx - tlx)
            x0, y0 = cx - w / 2, cy - h / 2
            if box is not None:
                rr = max(w, h)
                if cx + rr < box[0] or cx - rr > box[2] or cy + rr < box[1] or cy - rr > box[3]:
                    return
            if abs(ang) > 1e-4:
                rot = f' rot="{int(round(math.degrees(ang) * 60000)) % 21600000}"'
        else:
            x0, y0, w, h = o["x"], o["y"], o["w"], o["h"]
            if box is not None:
                cx0, cy0, cx1, cy1 = max(x0, box[0]), max(y0, box[1]), min(x0 + w, box[2]), min(y0 + h, box[3])
                if cx0 >= cx1 or cy0 >= cy1:
                    return
                if (cx0, cy0, cx1, cy1) != (x0, y0, x0 + w, y0 + h):          # кадрирование вместо обрезки
                    pct = lambda v, s: int(round(v / s * 100000))
                    crop = (f'<a:srcRect l="{pct(cx0 - x0, w)}" t="{pct(cy0 - y0, h)}" '
                            f'r="{pct(x0 + w - cx1, w)}" b="{pct(y0 + h - cy1, h)}"/>')
                    x0, y0, w, h = cx0, cy0, cx1 - cx0, cy1 - cy0
        x, y, ew, eh = self.e(x0), self.e(y0), max(1, self.e(w)), max(1, self.e(h))
        self.add(f'<p:pic><p:nvPicPr><p:cNvPr id="{self.nid()}" name="Рисунок"/><p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>'
                 f'<p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="{rid}"/>{crop}<a:stretch><a:fillRect/></a:stretch></p:blipFill>'
                 f'<p:spPr><a:xfrm{rot}><a:off x="{x}" y="{y}"/><a:ext cx="{ew}" cy="{eh}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
                 f'</p:spPr></p:pic>', (x, y, x + ew, y + eh))

    # --- обрезка ---
    def clip(self, o):
        self.flush()
        if o.get("rings"):
            rings = [r for r in o["rings"] if len(r) >= 3]
            self.clips.append({"box": _meet(self.box(), _bbox(rings)) if rings else self.box(), "rings": rings,
                               "items": self.stack[-1], "start": len(self.stack[-1])})
        else:
            x, y, w, h = o["rect"]
            self.clips.append({"box": _meet(self.box(), (x, y, x + w, y + h))})

    def unclip(self):
        self.flush()
        if not self.clips:
            return
        c = self.clips.pop()
        if not c.get("rings") or c["box"] is None or len(c["items"]) <= c["start"]:
            return
        # белая маска поверх нарисованного внутри контура: прямоугольник с вырезом по контуру
        x0, y0, x1, y1 = c["box"]
        k = EMU * self.k
        drawn = c["items"][c["start"]:]
        x0 = max(x0, min(b[0] for _, b in drawn) / k)
        y0 = max(y0, min(b[1] for _, b in drawn) / k)
        x1 = min(x1, max(b[2] for _, b in drawn) / k)
        y1 = min(y1, max(b[3] for _, b in drawn) / k)
        if x0 >= x1 or y0 >= y1:
            return
        pad = 0.01
        frame = [(x0 - pad, y0 - pad), (x1 + pad, y0 - pad), (x1 + pad, y1 + pad), (x0 - pad, y1 + pad)]
        outer = self.clips
        self.clips = []                                             # маска сама по себе не режется
        self._shape(orient([frame] + [list(r) for r in c["rings"]]), True, (1, 1, 1), None, 0, None, None, "Маска")
        self.clips = outer

    def body(self):
        """Фигуры слайда — разметка детей p:spTree."""
        self.flush()
        return "".join(x for x, _ in self.stack[0])


class PptxBackend:
    """pages: [(w_mm, h_mm, [(имя группы, ops)])] → bytes PPTX. Размер слайда — по первой странице, остальные
    вписываются в него с сохранением пропорций."""

    def __init__(self, title="Альбом"):
        self.title = title
        self._fonts, self._media, self._used = {}, {}, set()

    def font(self, family):
        self._used.add(family)                                      # какие начертания реально в тексте: их и встраиваем
        if family not in self._fonts:
            self._fonts[family] = Font(find_font(family))
        return self._fonts[family]

    def media(self, data):
        key = hashlib.sha1(data).hexdigest()
        self._media.setdefault(key, data)
        return key

    def _page(self, w, h, groups, W, H):
        s = _Slide(self, min(W / w, H / h))
        for name, ops in groups:
            s.open(name)
            depth = 0                                               # подгруппы слоёв внутри рамки
            for o in ops:
                t = o["op"]
                if t == "layer":
                    for _ in range(depth):
                        s.close()
                    s.open(o.get("name") or "Слой")
                    depth = 1
                elif t == "path":
                    s.alpha = o.get("alpha")
                    s.path(o["rings"], o.get("closed"), o.get("fill"), o.get("stroke"), o.get("width", 0.3), o.get("dash"), o.get("join"))
                    s.alpha = None
                elif t == "circle":
                    s.circle(o)
                elif t == "text":
                    s.text(o)
                elif t == "image":
                    s.image(o)
                elif t == "clip":
                    s.clip(o)
                elif t == "unclip":
                    s.unclip()
            for _ in range(depth):
                s.close()
            s.clips = []
            s.close()
        return s

    def _embed_fonts(self, prs):
        """Шрифты, использованные в тексте, встраиваются в файл (p:embeddedFontLst, части ppt/fonts/fontN.fntdata): без этого
        PowerPoint и Keynote на компьютере без Montserrat/PT подставляют другой шрифт и вид ломается. Как у Google Slides,
        часть — сам TrueType-файл в обёртке EOT без сжатия и шифрования (GUID-обфускация — схема Word, odttf; в PPTX её нет).
        python-pptx встраивание шрифтов не умеет: части и список добавляются через его пакет."""
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from pptx.opc.package import Part
        from pptx.opc.packuri import PackURI
        from pptx.oxml import parse_xml
        faces = {}                                                  # гарнитура → {"regular"|"bold": путь файла}
        for fam in sorted(self._used):
            if fam in FONTS and FONTS[fam][0] != "Courier New":
                face, bold = FONTS[fam]
                faces.setdefault(face, {})["bold" if bold else "regular"] = find_font(fam)
        if not faces:
            return
        lst, n = "", 0
        for face, kinds in faces.items():
            lst += f'<p:embeddedFont><p:font typeface="{_t(face)}"/>'
            for kind in ("regular", "bold"):
                if kind in kinds:
                    n += 1
                    part = Part(PackURI(f"/ppt/fonts/font{n}.fntdata"), "application/x-fontdata", prs.part.package,
                                _eot(Path(kinds[kind]).read_bytes()))
                    lst += f'<p:{kind} r:id="{prs.part.relate_to(part, RT.FONT)}"/>'
            lst += "</p:embeddedFont>"
        pres = prs.part._element
        pres.set("embedTrueTypeFonts", "1")
        pres.find(_P + "notesSz").addnext(parse_xml(f"<p:embeddedFontLst {NS}>{lst}</p:embeddedFontLst>"))

    def render(self, pages, notes=None):
        """notes — {номер слайда: текст заметки} (растровая карта на слайде)."""
        from pptx import Presentation
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from pptx.opc.packuri import PackURI
        from pptx.oxml import parse_xml
        from pptx.util import Emu
        W, H = pages[0][0], pages[0][1]
        prs = Presentation()                                        # стандартный пустой шаблон python-pptx
        prs.slide_width, prs.slide_height = Emu(int(round(W * EMU))), Emu(int(round(H * EMU)))
        prs.part._element.sldSz.attrib.pop("type", None)           # размер свой, не «экран 4:3»
        prs.slide_master.part.part_related_by(RT.THEME)._blob = (XML + THEME).encode("utf-8")   # цвета и шрифты ГРАДО
        props = prs.core_properties
        props.title, props.author, props.last_modified_by, props.comments = self.title, "ГРАДО 2", "ГРАДО 2", ""
        blank = next(lay for lay in prs.slide_layouts if lay.name == "Blank")
        for lay in [lay for lay in prs.slide_layouts if lay is not blank]:   # один макет «Пустой слайд», как прежде
            prs.slide_layouts.remove(lay)
        blank.name = "Пустой слайд"
        blank.part.partname = PackURI("/ppt/slideLayouts/slideLayout1.xml")
        pkg = prs.part.package                                      # без эскиза пустого слайда Office (его показал бы Finder)
        for rid, rel in list(pkg._rels.items()):
            if rel.reltype == RT.THUMBNAIL:
                pkg._rels.pop(rid)
        for i, (w, h, groups) in enumerate(pages, start=1):
            s = self._page(w, h, groups, W, H)
            slide = prs.slides.add_slide(blank)
            body = s.body()
            for key, ph in s.rels.items():
                body = body.replace(f'r:embed="{ph}"', f'r:embed="{slide.part.get_or_add_image_part(io.BytesIO(self._media[key]))[1]}"')
            tree = slide.shapes._spTree
            for el in parse_xml(f"<p:spTree {NS}>{body}</p:spTree>"):
                tree.append(el)
            if notes and notes.get(i):
                slide.notes_slide.notes_text_frame.text = notes[i]
        master = [rid for rid, rel in prs.part.rels.items() if rel.reltype == RT.NOTES_MASTER]
        if master:                                                  # python-pptx связывает образец заметок, но не вносит в список
            prs.part._element.find(_P + "sldMasterIdLst").addnext(
                parse_xml(f'<p:notesMasterIdLst {NS}><p:notesMasterId r:id="{master[0]}"/></p:notesMasterIdLst>'))
        self._embed_fonts(prs)
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()


def _scheme(name, val):
    return f'<a:{name}><a:srgbClr val="{val}"/></a:{name}>'


THEME = ('<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="ГРАДО">'
         '<a:themeElements><a:clrScheme name="ГРАДО">'
         + "".join(_scheme(n, v) for n, v in (("dk1", "000000"), ("lt1", "FFFFFF"), ("dk2", "1F1F1F"), ("lt2", "F2F2F2"),
                                              ("accent1", "E93939"), ("accent2", "2B3CC4"), ("accent3", "72A36F"),
                                              ("accent4", "F5D58A"), ("accent5", "A9CFE5"), ("accent6", "8A8A8A"),
                                              ("hlink", "2B3CC4"), ("folHlink", "7F3F98")))
         + '</a:clrScheme><a:fontScheme name="ГРАДО">'
         '<a:majorFont><a:latin typeface="Montserrat"/><a:ea typeface=""/><a:cs typeface=""/></a:majorFont>'
         '<a:minorFont><a:latin typeface="PT Sans"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont></a:fontScheme>'
         '<a:fmtScheme name="ГРАДО"><a:fillStyleLst>' + '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>' * 3
         + '</a:fillStyleLst><a:lnStyleLst>' + '<a:ln w="6350"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln>' * 3
         + '</a:lnStyleLst><a:effectStyleLst>' + '<a:effectStyle><a:effectLst/></a:effectStyle>' * 3
         + '</a:effectStyleLst><a:bgFillStyleLst>' + '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>' * 3
         + '</a:bgFillStyleLst></a:fmtScheme></a:themeElements><a:objectDefaults/><a:extraClrSchemeLst/></a:theme>')
